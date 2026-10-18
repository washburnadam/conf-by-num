"""Readers for Adam's legacy Excel files (read-only; used for validation baselines).

Legacy capture layout (one column per talk):
  row 1 = year, row 2 = month, row 3 = speaker, row 4+ = paragraphs.
  In ConferenceCapture_Apr2026.xlsx row 4 is the teaser (not spoken) -> first_body_row=5.
  In the older AllConferencesChronological sheet there is no teaser row, but bylines and
  'Notes' sections sometimes appear in the paragraph rows -> treat with care (validation only).

Legacy histogram layout: row 1 headers 'Unique/Output List/Total', row 2 counts,
  row 3 '#','Word','Freq', data from row 4.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl


@dataclass
class LegacyTalk:
    column: int
    year: int
    month: int
    speaker: str
    summary: str
    paragraphs: list[str] = field(default_factory=list)


def read_capture(path: str | Path, sheet: str | None = None, first_body_row: int = 5,
                 summary_row: int | None = 4) -> list[LegacyTalk]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    cols: dict[int, list] = {}
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        for c_idx, val in enumerate(row, start=1):
            cols.setdefault(c_idx, []).append(val)
    talks = []
    for c_idx, vals in sorted(cols.items()):
        if not vals or vals[0] in (None, ""):
            continue
        get = lambda r: vals[r - 1] if len(vals) >= r else None  # noqa: E731
        paras = [str(v).replace("_x000D_", " ") for v in vals[first_body_row - 1:] if v not in (None, "")]
        talks.append(LegacyTalk(
            column=c_idx, year=int(get(1)), month=int(get(2)),
            speaker=str(get(3) or "").strip(),
            summary=str(get(summary_row) or "").strip() if summary_row else "",
            paragraphs=paras,
        ))
    wb.close()
    return talks


def read_histogram(path: str | Path) -> dict[str, int]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    out: dict[str, int] = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row and row[1] not in (None, ""):
            out[str(row[1])] = int(row[2])
    wb.close()
    return out
