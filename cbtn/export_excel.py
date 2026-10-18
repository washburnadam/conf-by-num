"""On-demand Excel exports in Adam's legacy layouts (written to outputs/, disposable).

histogram:  row1 'Unique','Output List','Total' / row2 counts / row3 '#','Word','Freq' / data
capture:    one column per talk; row1 year, row2 month, row3 speaker (bold), row4 summary
            (yellow + italic), row5+ spoken paragraphs
"""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .archive import iter_talks


def histogram_xlsx(hist_csv: str | Path, out: str | Path) -> Path:
    with open(hist_csv, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    wb = Workbook()
    ws = wb.active
    ws.title = "WordHistogram"
    ws.append(["Unique", "Output List", "Total"])
    ws.append([len(rows), None, sum(int(r["count"]) for r in rows)])
    ws.append(["#", "Word", "Freq", "Rate per 10k", "Talks using", "Speakers using"])
    for r in rows:
        ws.append([int(r["rank"]), r["word"], int(r["count"]), float(r["rate_per_10k"]),
                   int(r["talks_using"]), int(r["speakers_using"])])
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def capture_xlsx(archive_root: str | Path, conf_key: str, out: str | Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "AllConferencesChronological"
    yellow = PatternFill("solid", fgColor="FFFF00")
    col = 0
    for t in iter_talks(archive_root, conf_key):
        if not t.counts_toward_histograms:
            continue
        col += 1
        ws.cell(1, col, t.conference["year"])
        ws.cell(2, col, t.conference["month"])
        c = ws.cell(3, col, t.speaker.name)
        c.font = Font(bold=True)
        c = ws.cell(4, col, t.summary or None)
        c.fill, c.font = yellow, Font(italic=True)
        r = 5
        for p in t.spoken_paragraphs():
            ws.cell(r, col, p.spoken_text)
            r += 1
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
