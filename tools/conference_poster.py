"""Conference by the Numbers — one infographic per conference (Adam, 2026-10-13) → outputs/final/<conf>/

  conference_by_the_numbers_<conf>.png   phone-first tall poster: classic header, emphasized words + phrases
                                         (packed bubbles, largest to smallest, labels to the right), rising words +
                                         phrases (sparkline cards), unique words (tags with the speaker)
  conference_by_the_numbers_<conf>.pptx  the same poster as native, editable PowerPoint shapes
  conference_by_the_numbers_<conf>.html  the same poster as SVG; click a bubble, card or tag for details + talk links

Every element is laid out ONCE in inches (Canvas) and drawn by two backends (matplotlib PNG, python-pptx), so the
PNG and the slide match. Text widths are measured with the real Aharoni font file.

Selection
- Words: the plotted starting set (family_review.plain: top 300, >= 2x the 1971-to-last-conference average,
  minus names / possessives / everyday words / near-duplicates / Adam's drops; "Dallin H. Oaks" as one name).
- Phrases: from the targeted phrase list (config/phrases.csv), >= 2x usual and >= PHRASE_MIN standalone uses,
  people-name phrases left out (like names among words) except the merged prophet's name.
- Rising: family_review.continual_increase (five eras rising, last >= 2x first, steady trend), >= 1.5x usual now;
  phrases >= 5 uses; people's names left out.
- Unique: Adam's picks in config/unique_words_keep.yaml, each with the speaker who used it (surname).

Usage: python tools/conference_poster.py 2026-04 2026-10
"""
from __future__ import annotations

import csv
import glob
import math
import os
import sqlite3
import sys
from dataclasses import dataclass, field
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import yaml  # noqa: E402
from matplotlib.lines import Line2D  # noqa: E402
from matplotlib.patches import Circle, FancyBboxPatch, Polygon, Rectangle  # noqa: E402
from PIL import ImageFont  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
from cbtn.graphics import THEMES, conf_label  # noqa: E402
from family_review import MERGED_NAMES, continual_increase, load, plain, removed, series  # noqa: E402

T = THEMES["oct2025"]
OUT = ROOT / "outputs/final"
W = 4.6                      # poster width (in), phone-first: 1 pt ≈ 1.2 px on a 390 px phone screen
PHRASE_MIN = 10              # standalone uses for a targeted phrase (words: ~20, the 300th word's count)
HEAD = "Aharoni"             # display face (Office cloud font on Adam's PC)
BODY = "Segoe UI"            # small explanatory text
TOPIC_COLORS = ["#C8501E", "#0E82A3", "#D23B6B", "#20B29B", "#F36C0F", "#6A5E9C", "#12B4E3", "#8C8C8C"]
INK, MUTED, GRID = "#1E1E1E", "#5A5A5A", "#E4E4E4"

_AH = sorted(glob.glob(os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\FontCache\4\CloudFonts\Aharoni\*.ttf")))
_FONTS = {}


def text_w(s, size, face=HEAD, bold=False):
    """Width in inches of `s` at `size` pt in the real font (Aharoni file, else Segoe UI / Segoe UI Bold)."""
    key = (face, size, bold)
    if key not in _FONTS:
        path = (_AH[0] if face == HEAD and _AH else
                r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf")
        try:
            _FONTS[key] = ImageFont.truetype(path, size * 10)
        except OSError:
            _FONTS[key] = None
    f = _FONTS[key]
    return (f.getlength(s) / 10 / 72) if f else len(s) * size * 0.55 / 72


def wrap(s, size, width, face=BODY):
    lines, cur = [], ""
    for w in s.split():
        t = (cur + " " + w).strip()
        if text_w(t, size, face) <= width or not cur:
            cur = t
        else:
            lines.append(cur)
            cur = w
    return lines + ([cur] if cur else [])


# ---------------------------------------------------------------- canvas (inches, origin top-left)
@dataclass
class Canvas:
    h: float = 0.0
    ops: list = field(default_factory=list)
    bg: str = "#FFFFFF"
    hits: list = field(default_factory=list)       # clickable regions for the HTML version: (id, x0, y0, x1, y1)

    def hit(self, hid, x0, y0, x1, y1):
        self.hits.append((hid, x0, y0, x1, y1))

    def rect(self, x, y, w, h, fill, r=0.0, line=None):
        self.ops.append(("rect", x, y, w, h, fill, r, line))

    def circle(self, cx, cy, d, fill, line=None, lw=0.0):
        self.ops.append(("circle", cx, cy, d, fill, line, lw))

    def poly(self, pts, fill=None, line=None, lw=1.0, closed=True):
        self.ops.append(("poly", pts, fill, line, lw, closed))

    def line(self, x1, y1, x2, y2, color, lw=0.75):
        self.ops.append(("poly", [(x1, y1), (x2, y2)], None, color, lw, False))

    def text(self, x, y, s, size, color=INK, face=HEAD, ha="center", va="center", bold=None, italic=False):
        """One line of text anchored at (x, y)."""
        self.ops.append(("text", x, y, s, size, color, face, ha, va, face == HEAD if bold is None else bold, italic))

    def para(self, x, y, lines, size, color=INK, face=BODY, ha="left", lead=1.3, bold=False):
        """Lines from top y; returns the y below the last line."""
        for ln in lines:
            self.text(x, y + size * lead / 72 / 2, ln, size, color, face, ha, "center", bold)
            y += size * lead / 72
        return y


# ---------------------------------------------------------------- backends
def render_png(c: Canvas, out):
    fig = plt.figure(figsize=(W, c.h), dpi=200, facecolor=c.bg)
    tr = fig.dpi_scale_trans
    H = c.h
    for z, op in enumerate(c.ops):          # draw strictly in op order (lines would otherwise sit above shapes)
        k = op[0]
        z = 1 + z * 1e-4
        if k == "rect":
            _, x, y, w, h, fill, r, line = op
            if r:
                p = FancyBboxPatch((x, H - y - h), w, h, boxstyle=f"round,pad=0,rounding_size={r}", transform=tr,
                                   facecolor=fill or "none", edgecolor=line or "none", lw=0.8)
            else:
                p = Rectangle((x, H - y - h), w, h, transform=tr, facecolor=fill or "none", edgecolor=line or "none")
            p.set_zorder(z)
            fig.patches.append(p)
        elif k == "circle":
            _, cx, cy, d, fill, line, lw = op
            fig.patches.append(Circle((cx, H - cy), d / 2, transform=tr, facecolor=fill, edgecolor=line or "none",
                                      lw=lw, zorder=z))
        elif k == "poly":
            _, pts, fill, line, lw, closed = op
            xy = [(px, H - py) for px, py in pts]
            if fill or closed:
                fig.patches.append(Polygon(xy, closed=closed, transform=tr, facecolor=fill or "none",
                                           edgecolor=line or "none", lw=lw, zorder=z))
            else:
                fig.lines.append(Line2D([p[0] for p in xy], [p[1] for p in xy], transform=tr, color=line, lw=lw,
                                        solid_capstyle="round", zorder=z))
        elif k == "text":
            _, x, y, s, size, color, face, ha, va, bold, italic = op
            fam = [HEAD, BODY, "DejaVu Sans"] if face == HEAD else [BODY, "DejaVu Sans"]
            fig.text(x, H - y, s, transform=tr, fontsize=size, color=color, ha=ha,
                     va={"center": "center_baseline"}.get(va, va), fontfamily=fam,
                     fontweight="bold" if bold else "normal", fontstyle="italic" if italic else "normal", zorder=z)
    fig.savefig(out, facecolor=c.bg)
    plt.close(fig)


def render_pptx(canvases, out):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    def rgb(h):
        return RGBColor.from_string(h.lstrip("#").upper())

    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(max(c.h for c in canvases))
    for c in canvases:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(c.bg)
        for op in c.ops:
            k = op[0]
            if k == "rect":
                _, x, y, w, h, fill, r, line = op
                sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,
                                        Inches(x), Inches(y), Inches(w), Inches(h))
                if r:
                    sh.adjustments[0] = min(0.5, r / min(w, h))
                _style(sh, fill, line, 0.8, rgb)
            elif k == "circle":
                _, cx, cy, d, fill, line, lw = op
                sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
                _style(sh, fill, line, lw, rgb)
            elif k == "poly":
                _, pts, fill, line, lw, closed = op
                fb = s.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
                fb.add_line_segments([(Inches(px), Inches(py)) for px, py in pts[1:]], close=closed)
                sh = fb.convert_to_shape()
                _style(sh, fill, line, lw, rgb)
            elif k == "text":
                _, x, y, txt, size, color, face, ha, va, bold, italic = op
                w = text_w(txt, size, face) + 0.12
                h = size * 1.35 / 72
                bx = x - w / 2 if ha == "center" else (x - w if ha == "right" else x)
                by = y - h / 2 if va == "center" else (y - h if va in ("bottom", "baseline") else y)
                tb = s.shapes.add_textbox(Emu(int(Inches(bx))), Inches(by), Inches(w), Inches(h))
                tf = tb.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = False
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = tf.paragraphs[0]
                p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(ha, PP_ALIGN.LEFT)
                run = p.add_run()
                run.text = txt
                run.font.size = Pt(size)
                run.font.name = face
                run.font.bold = bool(bold) and face != HEAD   # Aharoni is already bold-only
                run.font.italic = italic
                run.font.color.rgb = rgb(color)
    prs.save(out)


def _style(sh, fill, line, lw, rgb):
    from pptx.util import Pt
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw or 0.75)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False


# ---------------------------------------------------------------- data
def emphasized_items(con, conf):
    cur, allw, new, disp, phr = load(conf)
    cats = {r["canonical"]: r["category"] for r in csv.DictReader(open(ROOT / "config/phrases.csv", encoding="utf-8"))}
    items = [{"key": r["word"], "label": disp.get(r["word"], r["word"]), "n": int(r["count"]), "x": float(r["ratio"]),
              "phrase": bool(r.get("phrase"))} for r in cur if plain(r)]
    have = {i["key"] for i in items}
    skipped = []
    for p in phr:
        if not p["ratio"] or float(p["ratio"]) < 2 or p["phrase"] in have:
            continue
        if int(p["count"]) < PHRASE_MIN:
            continue
        if cats.get(p["phrase"], "").startswith("People") and p["phrase"] not in MERGED_NAMES:
            skipped.append(p["phrase"])
            continue
        items.append({"key": p["phrase"], "label": p["phrase"], "n": int(p["count"]), "x": float(p["ratio"]),
                      "phrase": True})
    return items, skipped, cur, allw, phr, disp


def topic_of(items):
    topics = yaml.safe_load((ROOT / "config/word_topics.yaml").read_text(encoding="utf-8"))
    where = {w: t for t, ws in topics.items() for w in ws}
    rows = {}
    for i in items:
        rows.setdefault(where.get(i["key"], "Other"), []).append(i)
    return rows


# ---------------------------------------------------------------- rising + unique data
def rising_items(con, conf, cur, allw, phr, disp):
    names = {r["word"] for r in cur if "name" in r["flags"] or removed(r)}
    out = []
    for r in allw.values():
        if (r["lexicon"] == "False" and r["new_word"] == "False" and r["ratio"] and int(r["rank"]) <= 300
                and float(r["ratio"]) >= 1.5 and len(r["word"]) > 1 and r["word"] not in names
                and r["possessive_flag"] != "True"):
            _, rates, _ = series(con, conf, word=r["word"])
            if continual_increase(rates[:-1]):
                out.append({"label": disp.get(r["word"], r["word"]), "n": int(r["count"]), "x": float(r["ratio"]),
                            "rates": rates, "phrase": False})
    ids = dict((r[1], r[0]) for r in con.execute("SELECT phrase_id, canonical FROM phrases"))
    cats = {r["canonical"]: r["category"] for r in csv.DictReader(open(ROOT / "config/phrases.csv", encoding="utf-8"))}
    for p in phr:   # people's names left out (as among words); >= 5 uses so the trend line is not noise
        if (p["ratio"] and float(p["ratio"]) >= 1.5 and int(p["count"]) >= 5 and p["phrase"] in ids
                and not cats.get(p["phrase"], "").startswith("People")):
            _, rates, _ = series(con, conf, pid=ids[p["phrase"]])
            if continual_increase(rates[:-1]):
                out.append({"label": p["phrase"], "n": int(p["count"]), "x": float(p["ratio"]), "rates": rates,
                            "phrase": True})
    out.sort(key=lambda i: -i["x"])
    return out[:18]


def surname(s):
    s = s.replace(", Jr.", "").replace(" Jr.", "")
    return s.split()[-1]


def unique_items(con, conf, disp):
    keep = yaml.safe_load((ROOT / "config/unique_words_keep.yaml").read_text(encoding="utf-8")).get(conf, [])
    out = []
    for w in keep:
        t = con.execute("SELECT t.speaker_display FROM word_counts_raw c JOIN words w USING(word_id) JOIN talks t "
                        "USING(tid) WHERE w.word=? AND t.conf_key=? ORDER BY c.count DESC LIMIT 1",
                        (w.lower(), conf)).fetchone()
        out.append({"key": w.lower(), "word": disp.get(w.lower(), w) if w != "Lego" else "Lego",
                    "who": surname(t[0]) if t else ""})
    return out


# ---------------------------------------------------------------- shared pieces (phone-first: W = 4.6 in)
M = 0.22                     # side margin (in)


def header(c: Canvas, conf):
    y = 0.0
    c.rect(0, y, W, 1.0, T["banner"])
    c.text(W / 2, y + 0.33, "CONFERENCE…", 25, "#FFFFFF")
    c.text(W / 2, y + 0.72, "…BY THE NUMBERS", 25, "#FFFFFF")
    y += 1.0
    c.rect(0, y, W, 0.6, T["band"])
    c.text(W / 2, y + 0.2, f"{conf_label(conf)} General Conference of", 10, "#FFFFFF", BODY, bold=False)
    c.text(W / 2, y + 0.4, "The Church of Jesus Christ of Latter-day Saints", 10, "#FFFFFF", BODY, bold=False)
    return y + 0.6


def section_bar(c, y, title, sub):
    lines = wrap(sub, 8.5, W - 2 * M)
    h = 0.42 + 8.5 * 1.3 / 72 * len(lines) + 0.1
    c.rect(0, y, W, h, T["header2"])
    c.text(W / 2, y + 0.24, title, 15, "#FFFFFF")
    yy = y + 0.42
    for ln in lines:
        c.text(W / 2, yy + 0.07, ln, 8.5, "#FFFFFF", BODY, bold=False)
        yy += 8.5 * 1.3 / 72
    return y + h + 0.16


def footer(c: Canvas, y, note):
    c.line(M, y, W - M, y, GRID)
    for ln in wrap(note, 7.5, W - 2 * 0.5):
        c.text(W / 2, y + 0.17, ln, 7.5, MUTED, BODY, bold=False)
        y += 0.15
    return y + 0.45                                    # room so rounded-corner image viewers don't clip it


# ---------------------------------------------------------------- packed bubble chart
def _hit_cc(a, b, gap):
    return math.hypot(a[0] - b[0], a[1] - b[1]) < a[2] + b[2] + gap


def _hit_rr(a, b, gap):
    return a[0] < b[2] + gap and b[0] < a[2] + gap and a[1] < b[3] + gap and b[1] < a[3] + gap


def _hit_cr(c, r, gap):
    nx, ny = min(max(c[0], r[0]), r[2]), min(max(c[1], r[1]), r[3])
    return math.hypot(c[0] - nx, c[1] - ny) < c[2] + gap


def pack_bubbles(items, X, top, x_left, x_right, diam, LS):
    """Largest to smallest, top to bottom. Each circle starts half-overlapping the previous one vertically
    (its top at the previous centre) and moves down only until it and its label clear everything placed."""
    placed, prev_cy = [], None
    GAP = 0.025
    for it in sorted(items, key=lambda i: -i["n"]):
        d = diam(it["n"])
        r = d / 2
        cx = X(it["x"])
        num_inside = d >= 0.24
        lab = it["label"] + ("" if num_inside else f" ({it['n']})")
        face = BODY if it["phrase"] else HEAD
        lw, lh = text_w(lab, LS, face, bold=it["phrase"]) * 1.06, LS * 1.25 / 72   # phrases: bold italic
        sides = [s for s in (True, False) if (cx + r + 0.04 + lw <= x_right if s else cx - r - 0.04 - lw >= x_left)]
        cy = top + r if prev_cy is None else max(top + r, prev_cy + r)

        def shapes(cy, right):
            lx0 = cx + r + 0.04 if right else cx - r - 0.04 - lw
            return (cx, cy, r), (lx0, cy - lh / 2, lx0 + lw, cy + lh / 2)

        def clear(circ, rect):
            return not any(_hit_cc(circ, p["circ"], GAP) or _hit_cr(circ, p["rect"], GAP) or
                           _hit_cr(p["circ"], rect, GAP) or _hit_rr(rect, p["rect"], GAP) for p in placed)

        right = None
        while right is None:          # label on the right if it fits there, else on the left; else move down
            for side in sides:
                circ, rect = shapes(cy, side)
                if clear(circ, rect):
                    right = side
                    break
            else:
                cy += 0.01
        placed.append({"it": it, "circ": circ, "rect": rect, "d": d, "lab": lab, "face": face, "right": right,
                       "num": num_inside})
        prev_cy = cy
    return placed


def emphasized_section(c, y, items, conf):
    y = section_bar(c, y, "Emphasized Words and Phrases", "What did the speakers say much more than usual?")
    lines_w = wrap("Words: of the 300 most-used words in this conference, those spoken at least twice as often (per "
                   "10,000 words) as in the average general conference since 1971. Names and everyday words are left "
                   "out.", 8.5, W - 2 * M)
    lines_p = wrap("Phrases: some targeted phrases we track (not every possible phrase) that pass the same test, with "
                   f"at least {PHRASE_MIN} uses.", 8.5, W - 2 * M)
    y = c.para(M, y, lines_w, 8.5, INK)
    y = c.para(M, y + 0.04, lines_p, 8.5, INK) + 0.14
    # legend 1: the x axis
    ay = y + 0.17
    aw = W - 2 * M
    c.poly([(M, ay - 0.11), (M + aw - 0.32, ay - 0.11), (M + aw - 0.32, ay - 0.2), (M + aw, ay),
            (M + aw - 0.32, ay + 0.2), (M + aw - 0.32, ay + 0.11), (M, ay + 0.11)], fill=T["right"])
    c.text(M + (aw - 0.3) / 2, ay, "further right = used more than usual", 9.5, "#FFFFFF")
    y = ay + 0.32
    # legend 2: size, word/phrase, topic colors
    maxn = max(i["n"] for i in items)
    DMAX = 0.62

    def diam(n):
        return DMAX * math.sqrt(n / maxn)

    lx = M
    samples = [s for s in (10, 50, 200) if s <= maxn * 1.05] or [maxn]
    base = y + diam(samples[-1])
    for s in samples:
        d = diam(s)
        c.circle(lx + d / 2, base - d / 2, d, "#C4C4C4")
        c.text(lx + d / 2, base + 0.1, f"{s}", 8, MUTED, BODY, bold=False)
        lx += d + 0.14
    c.para(lx + 0.05, y + 0.02, wrap("bigger circle = used more times (area = uses)", 8.5, W - M - lx - 0.05), 8.5,
           MUTED)
    c.circle(lx + 0.13, base - 0.02, 0.17 - 1.2 / 72, "#FFFFFF", MUTED, 1.2)
    c.text(lx + 0.28, base - 0.02, "ring = phrase", 8.5, MUTED, BODY, ha="left", bold=False)
    y = base + 0.26
    topics = yaml.safe_load((ROOT / "config/word_topics.yaml").read_text(encoding="utf-8"))
    where = {w: t for t, ws in topics.items() for w in ws}
    used = []
    for i in sorted(items, key=lambda i: -i["n"]):
        t = where.get(i["key"], "Other")
        if t not in used:
            used.append(t)
    col_of = {t: TOPIC_COLORS[k % len(TOPIC_COLORS)] for k, t in enumerate(used)}
    x = M
    for t in used:                                     # color key: chips that wrap
        w = 0.16 + text_w(t, 8.5, BODY) + 0.18
        if x + w > W - M:
            x, y = M, y + 0.2
        c.circle(x + 0.06, y + 0.08, 0.12, col_of[t])
        c.text(x + 0.16, y + 0.08, t, 8.5, INK, BODY, ha="left", bold=False)
        x += w
    y += 0.36
    # chart
    x_left, x_right = M, W - M
    xmax = math.log2(max(i["x"] for i in items))
    lo, hi = 1 - 0.2, xmax + 0.25
    px0, px1 = M + 0.25, W - M - 0.25

    def X(v):
        return px0 + (math.log2(v) - lo) / (hi - lo) * (px1 - px0)

    ticks = [2 ** k for k in range(1, math.floor(hi) + 1)]
    for t in ticks:
        c.text(X(t), y, f"{t}×", 9, MUTED)
    top = y + 0.16
    LS = 8.5
    placed = pack_bubbles(items, X, top + 0.02, x_left, x_right, diam, LS)
    bottom = max(p["circ"][1] + p["circ"][2] for p in placed) + 0.08
    for t in ticks:
        c.line(X(t), top, X(t), bottom, GRID, 0.6)
    for p in placed:
        it, (cx, cy, r) = p["it"], p["circ"]
        col = col_of[where.get(it["key"], "Other")]
        if it["phrase"]:
            lw = 1.2                                       # pt; stroke straddles the path, so shrink by it
            c.circle(cx, cy, max(p["d"] - lw / 72, 0.02), "#FFFFFF", col, lw)
        else:
            c.circle(cx, cy, p["d"], col, "#FFFFFF", 0.6)
        c.hit("e:" + it["key"], min(cx - r, p["rect"][0]), cy - max(r, 0.08), max(cx + r, p["rect"][2]),
              cy + max(r, 0.08))
        if p["num"]:
            c.text(cx, cy, str(it["n"]), 7.5, col if it["phrase"] else "#FFFFFF", BODY, bold=True)
        lx = cx + r + 0.04 if p["right"] else cx - r - 0.04
        c.text(lx, cy, p["lab"], LS, INK, p["face"], ha="left" if p["right"] else "right",
               bold=None if not it["phrase"] else True, italic=it["phrase"])
    for t in ticks:
        c.text(X(t), bottom + 0.12, f"{t}×", 9, MUTED)
    c.text(W / 2, bottom + 0.3, "times as often as in the average past conference", 8.5, MUTED, BODY, bold=False)
    return bottom + 0.55


def rising_section(c, y, rising, conf):
    y = section_bar(c, y, "Rising Words and Phrases", "Used more than usual this time AND climbing steadily for "
                    "decades: a long-term trend, not a one-time spike")
    cols, gap = 2, 0.14
    cw = (W - 2 * M - gap) / cols
    ch = 0.98
    for k, it in enumerate(rising):
        col, row = k % cols, k // cols
        x0, y0 = M + col * (cw + gap), y + row * (ch + 0.1)
        color = T["band"] if it["phrase"] else T["banner"]
        c.rect(x0, y0, cw, ch, "#F7F5F1", r=0.07)
        c.hit("r:" + it["label"], x0, y0, x0 + cw, y0 + ch)
        size = 11 if text_w(it["label"], 11) <= cw - 0.16 else 9
        c.text(x0 + 0.08, y0 + 0.15, it["label"], size, color, HEAD, ha="left", italic=it["phrase"])
        c.text(x0 + 0.08, y0 + 0.32, f"{it['x']:.1f}× usual · {it['n']} uses", 8, MUTED, BODY, ha="left",
               bold=False)
        sx0, sx1, sy0, sy1 = x0 + 0.08, x0 + cw - 0.1, y0 + 0.43, y0 + ch - 0.2
        v = it["rates"]
        mx = max(v) or 1

        def sx(i):
            return sx0 + i / (len(v) - 1) * (sx1 - sx0)
        pts = [(sx(i), sy1 - val / mx * (sy1 - sy0)) for i, val in enumerate(v)]
        c.poly([(sx0, sy1)] + pts + [(sx1, sy1)], fill="#E6E1D8", line=None, closed=True)
        c.line(sx0, sy1, sx1, sy1, "#C9C2B6", 0.5)                 # time axis: a tick per decade,
        for yr in range(1980, int(conf[:4]) + 1, 10):              # labels every other decade
            i = (yr - 1971) * 2                                    # two conferences a year from 1971-04
            if i < len(v):
                c.line(sx(i), sy1, sx(i), sy1 + 0.035, "#9A9389", 0.5)
                if yr % 20 == 0:
                    c.text(sx(i), sy1 + 0.1, str(yr), 7, MUTED, BODY, bold=False)
        c.poly(pts, None, color, 0.8, closed=False)
        c.circle(pts[-1][0], pts[-1][1], 0.065, color)
    y += math.ceil(len(rising) / cols) * (ch + 0.1)
    y = c.para(M, y, wrap("Each line: uses per 10,000 words at every general conference, 1971 → " + conf[:4] +
                          " (ticks = decades; dot = this conference) · orange = word, blue = phrase", 8, W - 2 * M), 8, MUTED)
    return y + 0.25


def unique_section(c, y, uniq, conf):
    y = section_bar(c, y, f"Unique Words", f"Spoken in general conference for the first time (since 1971) in "
                    f"{conf_label(conf)}, with the speaker who said it")
    pal = [T["banner"], T["band"], T["point"], T["panel"], "#6A5E9C", T["up"]]
    x, row, ry = M, 0, y
    tag_h, gap = 0.48, 0.1
    for k, u in enumerate(uniq):
        tw = max(text_w(u["word"], 11), text_w(u["who"], 8, BODY)) + 0.28
        if x + tw > W - M:
            row += 1
            x = M + (0.18 if row % 2 else 0)               # stagger alternate rows
            ry = y + row * (tag_h + gap)
        c.rect(x, ry, tw, tag_h, pal[k % len(pal)], r=0.08)
        c.hit("u:" + u["key"], x, ry, x + tw, ry + tag_h)
        c.text(x + tw / 2, ry + 0.19, u["word"], 11, "#FFFFFF")
        c.text(x + tw / 2, ry + 0.37, u["who"], 8, "#FFFFFF", BODY, bold=False)
        x += tw + gap
    return y + (row + 1) * (tag_h + gap) + 0.12


# ---------------------------------------------------------------- interactive HTML (same layout, drawn as SVG)
def _detail(con, conf, key, phrase, n=5):
    if phrase:
        pid = con.execute("SELECT phrase_id FROM phrases WHERE canonical=?", (key,)).fetchone()[0]
        confs, rates, _ = series(con, conf, pid=pid)
        tk = con.execute("SELECT t.speaker_display, t.title, SUM(p.standalone) AS k, t.url FROM phrase_counts p "
                         "JOIN talks t USING(talk_id) WHERE p.phrase_id=? AND t.conf_key=? AND p.standalone>0 "
                         "GROUP BY t.talk_id ORDER BY k DESC LIMIT ?", (pid, conf, n)).fetchall()
    else:
        confs, rates, _ = series(con, conf, word=key)
        tk = con.execute("SELECT t.speaker_display, t.title, c.count, t.url FROM word_counts_raw c JOIN words w "
                         "USING(word_id) JOIN talks t USING(tid) WHERE w.word=? AND t.conf_key=? AND "
                         "t.item_type='talk' ORDER BY c.count DESC LIMIT ?", (key, conf, n)).fetchall()
    return confs, [round(v, 3) for v in rates], [{"sp": a, "ti": b.strip("“”\"' "), "c": k, "u": u}
                                                  for a, b, k, u in tk]


def build_info(con, conf, items, rising, uniq):
    topics = yaml.safe_load((ROOT / "config/word_topics.yaml").read_text(encoding="utf-8"))
    where = {w: t for t, ws in topics.items() for w in ws}
    info, confs = {}, []
    for it in items:
        confs, rates, talks = _detail(con, conf, it["key"], it["phrase"])
        info["e:" + it["key"]] = {"kind": "emphasized", "label": it["label"], "phrase": it["phrase"], "n": it["n"],
                                  "x": round(it["x"], 1), "topic": where.get(it["key"], "Other"), "rates": rates,
                                  "talks": talks}
    for it in rising:
        key = it["label"] if it["phrase"] else it["label"].lower()
        _, _, talks = _detail(con, conf, key, it["phrase"])
        info["r:" + it["label"]] = {"kind": "rising", "label": it["label"], "phrase": it["phrase"], "n": it["n"],
                                    "x": round(it["x"], 1), "rates": [round(v, 3) for v in it["rates"]],
                                    "talks": talks}
    for u in uniq:
        _, _, talks = _detail(con, conf, u["key"], False, 3)
        info["u:" + u["key"]] = {"kind": "unique", "label": u["word"], "talks": talks,
                                 "n": sum(t["c"] for t in talks)}
    return info, confs


def _xml(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def render_svg(c: Canvas):
    P = 72.0
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W * P:.1f} {c.h * P:.1f}" role="img" '
           f'aria-label="Conference by the Numbers infographic">', f'<rect width="100%" height="100%" fill="{c.bg}"/>']
    for op in c.ops:
        k = op[0]
        if k == "rect":
            _, x, y, w, h, fill, r, line = op
            out.append(f'<rect x="{x * P:.2f}" y="{y * P:.2f}" width="{w * P:.2f}" height="{h * P:.2f}" '
                       f'rx="{r * P:.2f}" fill="{fill or "none"}" stroke="{line or "none"}"/>')
        elif k == "circle":
            _, cx, cy, d, fill, line, lw = op
            out.append(f'<circle cx="{cx * P:.2f}" cy="{cy * P:.2f}" r="{d / 2 * P:.2f}" fill="{fill}" '
                       f'stroke="{line or "none"}" stroke-width="{lw or 0}"/>')
        elif k == "poly":
            _, pts, fill, line, lw, closed = op
            pp = " ".join(f"{px * P:.2f},{py * P:.2f}" for px, py in pts)
            tag = "polygon" if closed else "polyline"
            out.append(f'<{tag} points="{pp}" fill="{fill or "none"}" stroke="{line or "none"}" '
                       f'stroke-width="{lw}" stroke-linejoin="round" stroke-linecap="round"/>')
        elif k == "text":
            _, x, y, txt, size, color, face, ha, va, bold, italic = op
            anchor = {"left": "start", "center": "middle", "right": "end"}[ha]
            fam = "head" if face == HEAD else "body"
            tl = text_w(txt, size, face, bold=bool(bold) and face != HEAD) * P
            out.append(f'<text class="{fam}{" b" if bold and face != HEAD else ""}{" i" if italic else ""}" '
                       f'x="{x * P:.2f}" y="{y * P:.2f}" font-size="{size}" fill="{color}" '
                       f'text-anchor="{anchor}" dominant-baseline="central" textLength="{tl:.1f}" '
                       f'lengthAdjust="spacingAndGlyphs">{_xml(txt)}</text>')
    for hid, x0, y0, x1, y1 in c.hits:
        out.append(f'<rect class="hot" data-id="{_xml(hid)}" x="{(x0 - 0.03) * P:.2f}" y="{(y0 - 0.03) * P:.2f}" '
                   f'width="{(x1 - x0 + 0.06) * P:.2f}" height="{(y1 - y0 + 0.06) * P:.2f}" rx="6" tabindex="0" '
                   f'role="button"/>')
    out.append("</svg>")
    return "\n".join(out)


def render_html(c: Canvas, info, confs, conf, out):
    import json
    tpl = (ROOT / "tools/poster_template.html").read_text(encoding="utf-8")
    data = {"label": conf_label(conf), "confs": confs, "info": info}
    html = (tpl.replace("{{TITLE}}", f"Conference by the Numbers · {conf_label(conf)}")
            .replace("<!--SVG-->", render_svg(c)).replace("/*DATA*/null", json.dumps(data, ensure_ascii=False)))
    out.write_text(html, encoding="utf-8")


def poster(con, conf):
    items, skipped, cur, allw, phr, disp = emphasized_items(con, conf)
    rising = rising_items(con, conf, cur, allw, phr, disp)
    uniq = unique_items(con, conf, disp)
    c = Canvas()
    y = header(c, conf) + 0.0
    y = emphasized_section(c, y, items, conf)
    y = rising_section(c, y, rising, conf)
    y = unique_section(c, y, uniq, conf)
    c.h = footer(c, y, "Counts are of the spoken text of every general conference talk, 1971–" + conf[:4] +
                 " · Conference by the Numbers")
    return c, items, skipped, rising, uniq


def main(confs):
    con = sqlite3.connect(ROOT / "data/cbtn.sqlite")
    for conf in confs:
        out = OUT / conf
        out.mkdir(parents=True, exist_ok=True)
        c, items, skipped, rising, uniq = poster(con, conf)
        render_png(c, out / f"conference_by_the_numbers_{conf}.png")
        render_pptx([c], out / f"conference_by_the_numbers_{conf}.pptx")
        info, confs = build_info(con, conf, items, rising, uniq)
        render_html(c, info, confs, conf, out / f"conference_by_the_numbers_{conf}.html")
        ph = sum(i["phrase"] for i in items)
        print(f"{conf}: {len(items) - ph} words + {ph} phrases (people phrases left out: {skipped}); rising "
              f"{len(rising)}; unique {len(uniq)}; poster {W} x {c.h:.1f} in")


if __name__ == "__main__":
    main(sys.argv[1:])
